#!/usr/bin/env python3
"""
validate_doc.py — Multi-format document validator.

Usage:
    # Validate a single file
    python scripts/validate_doc.py --file output/report.pdf

    # Validate all files in a directory
    python scripts/validate_doc.py --dir output/

    # Smoke test: generate one demo file per format, then validate all
    python scripts/validate_doc.py --smoke-test
"""
import argparse
import pathlib
import struct
import sys
import zipfile
from dataclasses import dataclass, field
from typing import Any


@dataclass
class ValidationResult:
    path: str
    ok: bool = True
    checks: list[str] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def fail(self, msg: str) -> None:
        self.ok = False
        self.issues.append(f"❌ {msg}")

    def warn(self, msg: str) -> None:
        self.warnings.append(f"⚠️  {msg}")

    def ok_check(self, msg: str) -> None:
        self.checks.append(f"  ✅ {msg}")

    def summary(self) -> str:
        status = "✅ PASS" if self.ok else "❌ FAIL"
        lines = [f"\n{'─'*60}", f"  {status}  {self.path}", f"{'─'*60}"]
        lines += self.checks + self.issues + self.warnings
        return "\n".join(lines)


# ── Magic byte signatures ─────────────────────────────────────────────────────
MAGIC: dict[str, bytes] = {
    ".pdf":  b"%PDF",
    ".docx": b"PK\x03\x04",  # ZIP container
    ".pptx": b"PK\x03\x04",  # ZIP container
    ".xlsx": b"PK\x03\x04",  # ZIP container
    ".rtf":  b"{\\rtf",
    ".html": None,  # no magic — check content
    ".md":   None,
    ".txt":  None,
}

OOXML_CONTENT_TYPES: dict[str, str] = {
    ".docx": "word/",
    ".pptx": "ppt/",
    ".xlsx": "xl/",
}


def _check_exists_and_size(r: ValidationResult, path: pathlib.Path) -> bool:
    if not path.exists():
        r.fail(f"File does not exist: {path}")
        return False
    r.ok_check(f"File exists: {path.name}")
    size = path.stat().st_size
    if size == 0:
        r.fail("File is zero bytes — generation failed")
        return False
    r.ok_check(f"File size: {size:,} bytes")
    return True


def _check_magic(r: ValidationResult, path: pathlib.Path) -> bool:
    ext = path.suffix.lower()
    expected = MAGIC.get(ext)
    if expected is None:
        return True  # no magic check for text formats
    try:
        header = path.read_bytes()[:8]
        if not header.startswith(expected):
            r.fail(f"Magic bytes mismatch for {ext} — file may be corrupt")
            return False
        r.ok_check(f"Magic bytes valid ({ext})")
    except OSError as e:
        r.fail(f"Cannot read file: {e}")
        return False
    return True


def _check_pdf(r: ValidationResult, path: pathlib.Path) -> None:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        page_count = len(reader.pages)
        if page_count == 0:
            r.fail("PDF has 0 pages")
        else:
            r.ok_check(f"PDF pages: {page_count}")

        # Try extracting text from first page
        text = reader.pages[0].extract_text() or ""
        if len(text.strip()) < 5:
            r.warn("First page appears to have no extractable text (may be image-only)")
        else:
            r.ok_check(f"Text extractable ({len(text)} chars on page 1)")

        # Check metadata
        meta = reader.metadata
        if meta:
            r.ok_check(f"Metadata present (title: {meta.get('/Title', 'none')})")

    except ImportError:
        r.warn("pypdf not installed — skipping PDF deep validation")
    except Exception as e:
        r.fail(f"PDF validation error: {e}")


def _check_ooxml(r: ValidationResult, path: pathlib.Path, fmt: str) -> None:
    """Validate DOCX / PPTX / XLSX (all are ZIP archives with OOXML content)."""
    expected_dir = OOXML_CONTENT_TYPES.get(path.suffix.lower(), "")
    try:
        with zipfile.ZipFile(str(path), "r") as z:
            names = z.namelist()
            # Check for required OOXML directory
            if expected_dir and not any(n.startswith(expected_dir) for n in names):
                r.fail(f"Missing OOXML content directory '{expected_dir}' in ZIP")
            else:
                r.ok_check(f"OOXML structure valid ({len(names)} entries in ZIP)")

            # Check [Content_Types].xml
            if "[Content_Types].xml" in names:
                r.ok_check("[Content_Types].xml present")
            else:
                r.fail("Missing [Content_Types].xml — file is not valid OOXML")

    except zipfile.BadZipFile:
        r.fail(f"{fmt} is not a valid ZIP/OOXML archive — file is corrupt")
    except Exception as e:
        r.fail(f"{fmt} validation error: {e}")


def _check_docx(r: ValidationResult, path: pathlib.Path) -> None:
    _check_ooxml(r, path, "DOCX")
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
        para_count = len(doc.paragraphs)
        if para_count == 0:
            r.warn("DOCX has no paragraphs")
        else:
            r.ok_check(f"DOCX paragraphs: {para_count}")
        r.ok_check(f"DOCX opens without error")
    except ImportError:
        r.warn("python-docx not installed — skipping DOCX deep validation")
    except Exception as e:
        r.fail(f"DOCX open error: {e}")


def _check_pptx(r: ValidationResult, path: pathlib.Path) -> None:
    _check_ooxml(r, path, "PPTX")
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        slide_count = len(prs.slides)
        if slide_count == 0:
            r.fail("PPTX has 0 slides")
        else:
            r.ok_check(f"PPTX slides: {slide_count}")
        r.ok_check("PPTX opens without error")
    except ImportError:
        r.warn("python-pptx not installed — skipping PPTX deep validation")
    except Exception as e:
        r.fail(f"PPTX open error: {e}")


def _check_xlsx(r: ValidationResult, path: pathlib.Path) -> None:
    _check_ooxml(r, path, "XLSX")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True)
        sheet_count = len(wb.sheetnames)
        if sheet_count == 0:
            r.fail("XLSX has no sheets")
        else:
            r.ok_check(f"XLSX sheets: {sheet_count}  ({', '.join(wb.sheetnames)})")
        wb.close()
        r.ok_check("XLSX opens without error")
    except ImportError:
        r.warn("openpyxl not installed — skipping XLSX deep validation")
    except Exception as e:
        r.fail(f"XLSX open error: {e}")


def _check_rtf(r: ValidationResult, path: pathlib.Path) -> None:
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
        if not content.startswith("{\\rtf"):
            r.fail("RTF does not start with {\\rtf — invalid format")
        elif not content.rstrip().endswith("}"):
            r.warn("RTF may be missing closing brace")
        else:
            r.ok_check("RTF header and footer valid")
            r.ok_check(f"RTF size: {len(content):,} chars")
    except Exception as e:
        r.fail(f"RTF read error: {e}")


def _check_html(r: ValidationResult, path: pathlib.Path) -> None:
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
        if "<html" not in content.lower():
            r.warn("HTML may be missing <html> tag")
        else:
            r.ok_check("HTML structure: <html> found")
        if "<body" not in content.lower():
            r.warn("HTML may be missing <body> tag")
        else:
            r.ok_check("HTML structure: <body> found")
        r.ok_check(f"HTML size: {len(content):,} chars")
    except Exception as e:
        r.fail(f"HTML read error: {e}")


def _check_text(r: ValidationResult, path: pathlib.Path) -> None:
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
        if len(content.strip()) < 5:
            r.warn("Text file appears nearly empty")
        else:
            r.ok_check(f"Text file: {len(content):,} chars, "
                       f"{content.count(chr(10))} lines")
    except Exception as e:
        r.fail(f"Text read error: {e}")


CHECKERS = {
    ".pdf":  _check_pdf,
    ".docx": _check_docx,
    ".pptx": _check_pptx,
    ".xlsx": _check_xlsx,
    ".rtf":  _check_rtf,
    ".html": _check_html,
    ".htm":  _check_html,
    ".md":   _check_text,
    ".txt":  _check_text,
}


def validate_file(path: str) -> ValidationResult:
    p = pathlib.Path(path)
    r = ValidationResult(path=str(p))

    if not _check_exists_and_size(r, p):
        return r
    if not _check_magic(r, p):
        return r

    ext = p.suffix.lower()
    checker = CHECKERS.get(ext)
    if checker:
        checker(r, p)
    else:
        r.warn(f"No deep validator for extension '{ext}' — basic checks only")

    return r


def validate_directory(dir_path: str) -> list[ValidationResult]:
    p = pathlib.Path(dir_path)
    if not p.is_dir():
        print(f"❌ Not a directory: {dir_path}", file=sys.stderr)
        return []
    results = []
    for f in sorted(p.iterdir()):
        if f.suffix.lower() in CHECKERS:
            results.append(validate_file(str(f)))
    return results


def smoke_test() -> list[ValidationResult]:
    """Generate one demo file per format, then validate each."""
    print("\n🔥 Smoke test — generating demo files for all formats...")
    scripts_dir = pathlib.Path(__file__).parent

    import subprocess

    tests = [
        ("generate_pdf.py",  "output/smoke_test.pdf",  "--title", "Smoke Test PDF"),
        ("generate_docx.py", "output/smoke_test.docx", "--title", "Smoke Test DOCX"),
        ("generate_pptx.py", "output/smoke_test.pptx", "--title", "Smoke Test PPTX"),
        ("generate_xlsx.py", "output/smoke_test.xlsx", "--title", "Smoke Test XLSX"),
        ("generate_rtf.py",  "output/smoke_test.rtf",  "--title", "Smoke Test RTF"),
    ]

    results = []
    for script, out_path, *extra_args in tests:
        script_path = scripts_dir / script
        if not script_path.exists():
            print(f"  ⚠️  Script not found: {script_path}")
            continue
        cmd = [sys.executable, str(script_path), "--output", out_path] + list(extra_args)
        env = {**__import__("os").environ, "PYTHONUTF8": "1"}
        proc = subprocess.run(cmd, capture_output=True, text=True, env=env)
        if proc.returncode != 0:
            print(f"  ❌ {script} failed:\n{proc.stderr}")
        else:
            print(f"  ✅ {script} → {out_path}")
        results.append(validate_file(out_path))

    return results


def print_results(results: list[ValidationResult]) -> int:
    """Print validation results. Returns exit code (0=all pass, 1=any fail)."""
    pass_count = sum(1 for r in results if r.ok)
    fail_count = len(results) - pass_count

    for r in results:
        print(r.summary())

    print(f"\n{'═'*60}")
    print(f"  TOTAL: {len(results)} file(s)  |  "
          f"PASS: {pass_count}  |  FAIL: {fail_count}")
    print(f"{'═'*60}\n")
    return 0 if fail_count == 0 else 1


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate generated document files.")
    parser.add_argument("--file",       help="Validate a single file")
    parser.add_argument("--dir",        help="Validate all supported files in a directory")
    parser.add_argument("--smoke-test", action="store_true",
                        help="Generate demo files for all formats, then validate")
    args = parser.parse_args()

    results: list[ValidationResult] = []

    if args.smoke_test:
        results = smoke_test()
    elif args.file:
        results = [validate_file(args.file)]
    elif args.dir:
        results = validate_directory(args.dir)
    else:
        parser.print_help()
        sys.exit(0)

    sys.exit(print_results(results))


if __name__ == "__main__":
    main()
