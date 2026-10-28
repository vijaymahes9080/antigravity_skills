#!/usr/bin/env python3
"""
generate_xlsx.py — Professional XLSX spreadsheet generator using openpyxl.

Usage:
    python scripts/generate_xlsx.py --title "Data Sheet" \
        --output output/data.xlsx [--brief brief.json]

Brief JSON format:
{
  "title": "Workbook Title",
  "author": "Author",
  "sheets": [
    {
      "name": "Sheet Name",
      "headers": ["Column A", "Column B", "Column C"],
      "rows": [["value1", "value2", "value3"]],
      "freeze_panes": true,
      "auto_filter": true
    }
  ]
}
"""
import argparse
import json
import pathlib
import sys
from datetime import datetime
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, GradientFill, PatternFill, Side
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    sys.exit("❌ openpyxl not installed. Run: python scripts/bootstrap.py")


# ─── Design System ───────────────────────────────────────────────────────────
NAVY_HEX    = "1A3380"
AMBER_HEX   = "F29910"
LIGHT_HEX   = "F5F7FA"
BORDER_HEX  = "C8CDD4"
WHITE_HEX   = "FFFFFF"
DARK_HEX    = "262626"

HEADER_FILL = PatternFill("solid", fgColor=NAVY_HEX)
ALT_FILL    = PatternFill("solid", fgColor=LIGHT_HEX)
HEADER_FONT = Font(name="Calibri", bold=True, color=WHITE_HEX, size=11)
BODY_FONT   = Font(name="Calibri", size=10, color=DARK_HEX)

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_HEX),
    right=Side(style="thin", color=BORDER_HEX),
    top=Side(style="thin", color=BORDER_HEX),
    bottom=Side(style="thin", color=BORDER_HEX),
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN   = Alignment(horizontal="left",   vertical="center", wrap_text=True)


def safe_cell_value(value: Any) -> Any:
    """Prevent formula injection in string cells."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _build_sheet(ws, sheet_data: dict[str, Any]) -> None:
    headers     = sheet_data.get("headers", [])
    rows        = sheet_data.get("rows", [])
    freeze      = sheet_data.get("freeze_panes", True)
    auto_filter = sheet_data.get("auto_filter", True)

    if not headers and not rows:
        return

    col_count = len(headers) if headers else (len(rows[0]) if rows else 1)

    # ── Header row ────────────────────────────────────────────────────────
    if headers:
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=str(header))
            cell.fill   = HEADER_FILL
            cell.font   = HEADER_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_ALIGN
        ws.row_dimensions[1].height = 22

    # ── Data rows ─────────────────────────────────────────────────────────
    data_start_row = 2 if headers else 1
    for row_idx, row_data in enumerate(rows, start=data_start_row):
        for col_idx in range(1, col_count + 1):
            val = row_data[col_idx - 1] if col_idx - 1 < len(row_data) else None
            cell = ws.cell(row=row_idx, column=col_idx,
                           value=safe_cell_value(val))
            # Alternating row fill
            cell.fill   = ALT_FILL if (row_idx % 2 == 0) else PatternFill("solid", fgColor=WHITE_HEX)
            cell.font   = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = LEFT_ALIGN

    # ── Auto-width columns ────────────────────────────────────────────────
    for col_idx in range(1, col_count + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    # ── Freeze top row ────────────────────────────────────────────────────
    if freeze and headers:
        ws.freeze_panes = ws.cell(row=2, column=1)

    # ── Auto-filter ───────────────────────────────────────────────────────
    if auto_filter and headers:
        last_col = get_column_letter(col_count)
        last_row = data_start_row + len(rows) - 1
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"


def generate(brief: dict[str, Any], output_path: str) -> str:
    title  = brief.get("title", "Workbook")
    author = brief.get("author", "")
    sheets = brief.get("sheets", [])

    out = pathlib.Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp.xlsx")

    wb = Workbook()
    wb.remove(wb.active)  # Remove default empty sheet

    wb.properties.title   = title
    wb.properties.creator = author or "document-maker-all-format"
    wb.properties.created = datetime.now()

    for sheet_data in sheets:
        name = sheet_data.get("name", f"Sheet{len(wb.sheetnames)+1}")[:31]
        ws = wb.create_sheet(title=name)
        _build_sheet(ws, sheet_data)

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")

    wb.save(str(tmp))
    tmp.replace(out)

    size_kb = out.stat().st_size // 1024
    sheet_count = len(wb.sheetnames)
    print(f"✅ XLSX generated: {out}  ({sheet_count} sheet(s), {size_kb} KB)")
    return str(out)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a professional XLSX spreadsheet.")
    parser.add_argument("--title",  default="Data Sheet")
    parser.add_argument("--author", default="")
    parser.add_argument("--output", default="output/data.xlsx")
    parser.add_argument("--brief",  default=None)
    args = parser.parse_args()

    if args.brief:
        brief = json.loads(pathlib.Path(args.brief).read_text(encoding="utf-8"))
    else:
        brief = {
            "title": args.title, "author": args.author,
            "sheets": [
                {
                    "name": "Format Support",
                    "headers": ["Format", "Engine", "Version", "Status", "Notes"],
                    "rows": [
                        ["PDF",      "reportlab",   "5.0.1",  "Active", "Full layout control"],
                        ["DOCX",     "python-docx", "1.2.0",  "Active", "Native OOXML"],
                        ["PPTX",     "python-pptx", "1.0+",   "Active", "16:9 widescreen"],
                        ["XLSX",     "openpyxl",    "3.1.5",  "Active", "Formulas, charts"],
                        ["RTF",      "Custom",      "1.0",    "Active", "Basic text & tables"],
                        ["HTML",     "Jinja2+lxml", "3.1.6",  "Active", "Semantic, accessible"],
                        ["Markdown", "markdown-it", "4.2.0",  "Active", "GFM compatible"],
                        ["Text",     "stdlib",      "builtin","Active", "Plain text, textwrap"],
                    ],
                    "freeze_panes": True,
                    "auto_filter": True,
                },
                {
                    "name": "Design Tokens",
                    "headers": ["Token", "Value", "Usage"],
                    "rows": [
                        ["color_primary",  "#1A3380", "Headings, header bars"],
                        ["color_accent",   "#F29910", "Dividers, highlights"],
                        ["color_body",     "#262626", "Body text"],
                        ["color_muted",    "#8C8C8C", "Captions, secondary text"],
                        ["font_heading",   "Calibri Bold / Helvetica-Bold", "H1–H3"],
                        ["font_body",      "Calibri / Helvetica", "Body paragraphs"],
                        ["size_h1",        "26–32pt", "Cover titles"],
                        ["size_h2",        "18–22pt", "Section headings"],
                        ["size_body",      "10–12pt", "Body text"],
                        ["margin_page",    "2–2.5cm", "All formats"],
                    ],
                    "freeze_panes": True,
                    "auto_filter": False,
                },
            ],
        }

    generate(brief, args.output)


if __name__ == "__main__":
    main()

